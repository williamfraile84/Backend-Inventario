import io
import os
from typing import Dict, Any, List, Optional
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

class ExportService:
    """Servicio de exportación formal de documentos a PDF y Excel."""

    @staticmethod
    def generate_invoice_pdf(invoice_data: Dict[str, Any], items: List[Dict[str, Any]]) -> bytes:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#065F46"),
            fontName="Helvetica-Bold",
        )
        meta_label = ParagraphStyle(
            "MetaLabel",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            fontName="Helvetica-Bold",
            textColor=colors.HexColor("#374151")
        )
        meta_val = ParagraphStyle(
            "MetaVal",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            fontName="Helvetica",
            textColor=colors.HexColor("#111827")
        )

        elements = []
        elements.append(Paragraph(f"FACTURA / COMPROBANTE - {invoice_data.get('provider_name') or 'FRUVER'}", title_style))
        elements.append(Spacer(1, 10))

        # Metadatos
        info_data = [
            [Paragraph("Proveedor:", meta_label), Paragraph(str(invoice_data.get("provider_name") or "N/A"), meta_val),
             Paragraph("No. Factura:", meta_label), Paragraph(str(invoice_data.get("invoice_number") or "N/A"), meta_val)],
            [Paragraph("NIT / Cédula:", meta_label), Paragraph(str(invoice_data.get("provider_nit") or "N/A"), meta_val),
             Paragraph("Fecha Emisión:", meta_label), Paragraph(str(invoice_data.get("invoice_date") or "N/A"), meta_val)],
        ]
        t_info = Table(info_data, colWidths=[80, 200, 90, 150])
        t_info.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_info)
        elements.append(Spacer(1, 15))

        # Tabla de Productos
        table_rows = [
            ["#", "Código", "Descripción", "Cant.", "Presentación", "Costo Unit.", "Subtotal"]
        ]
        for idx, it in enumerate(items):
            desc = it.get("description") or it.get("descripcion", "")
            cant = float(it.get("quantity") or 1.0)
            cost = float(it.get("unit_cost_base") or it.get("precio_unitario") or 0.0)
            sub = cant * cost
            pres = it.get("presentation") or it.get("unidad", "Und")
            code = it.get("product_code") or it.get("codigo") or ""
            table_rows.append([
                str(idx + 1),
                code[:15],
                desc[:35],
                f"{cant:g}",
                pres,
                f"${cost:,.0f}".replace(",", "."),
                f"${sub:,.0f}".replace(",", ".")
            ])

        t_prod = Table(table_rows, colWidths=[25, 75, 200, 45, 65, 65, 65])
        t_prod.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#065F46")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(t_prod)
        elements.append(Spacer(1, 15))

        # Totales
        tot_sub = float(invoice_data.get("subtotal") or 0.0)
        tot_tax = float(invoice_data.get("tax_total") or 0.0)
        tot_gen = float(invoice_data.get("total") or (tot_sub + tot_tax))
        totals_data = [
            ["Subtotal:", f"${tot_sub:,.0f}".replace(",", ".")],
            ["Impuestos (IVA/ICUI):", f"${tot_tax:,.0f}".replace(",", ".")],
            ["TOTAL A PAGAR:", f"${tot_gen:,.0f}".replace(",", ".")]
        ]
        t_tot = Table(totals_data, colWidths=[120, 100])
        t_tot.setStyle(TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("LINEABOVE", (0, 2), (-1, 2), 1, colors.HexColor("#065F46")),
        ]))
        elements.append(t_tot)

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_invoice_excel(invoice_data: Dict[str, Any], items: List[Dict[str, Any]]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Factura"

        header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB"),
        )

        ws["A1"] = f"FACTURA: {invoice_data.get('provider_name') or 'FRUVER'}"
        ws["A1"].font = Font(name="Arial", size=14, bold=True, color="065F46")

        ws["A3"] = "Proveedor:"
        ws["B3"] = invoice_data.get("provider_name") or ""
        ws["A4"] = "NIT:"
        ws["B4"] = invoice_data.get("provider_nit") or ""
        ws["D3"] = "No. Factura:"
        ws["E3"] = invoice_data.get("invoice_number") or ""
        ws["D4"] = "Fecha:"
        ws["E4"] = invoice_data.get("invoice_date") or ""

        for cell in ["A3", "A4", "D3", "D4"]:
            ws[cell].font = Font(bold=True)

        headers = ["Item", "Código", "Descripción", "Cantidad", "Presentación", "Costo Unitario", "Subtotal", "Margen %", "Precio Venta"]
        ws.append([])
        ws.append(headers)
        row_idx = 7

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx - 1, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        for idx, it in enumerate(items):
            desc = it.get("description") or it.get("descripcion", "")
            cant = float(it.get("quantity") or 1.0)
            cost = float(it.get("unit_cost_base") or it.get("precio_unitario") or 0.0)
            sub = cant * cost
            pres = it.get("presentation") or it.get("unidad", "Und")
            code = it.get("product_code") or it.get("codigo") or ""
            marg = float(it.get("margin_percent") or 30.0)
            sale = float(it.get("sale_price_final") or (cost * (1 + marg / 100.0)))

            row_data = [idx + 1, code, desc, cant, pres, cost, sub, marg, sale]
            ws.append(row_data)

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border_thin
                if col_idx in [6, 7, 9]:
                    cell.number_format = '"$"#,##0'
                elif col_idx == 4:
                    cell.number_format = '#,##0.00'
                elif col_idx == 8:
                    cell.number_format = '0.0"%"'
            row_idx += 1

        # Totales
        row_idx += 1
        ws.cell(row=row_idx, column=6, value="TOTAL:")
        ws.cell(row=row_idx, column=6).font = Font(bold=True)
        tot_cell = ws.cell(row=row_idx, column=7, value=f"=SUM(G7:G{row_idx-2})")
        tot_cell.font = Font(bold=True)
        tot_cell.number_format = '"$"#,##0'

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

